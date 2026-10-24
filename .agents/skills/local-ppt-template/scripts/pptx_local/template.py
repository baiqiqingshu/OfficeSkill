from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import shutil
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt

from .common import (PptError, identifier, json_bytes,
                     positive, read_json, require_keys, sha256)
from .package import clone_slides, inspect_package

EMU_CM = 360000


def blank_image() -> BytesIO:
    stream = BytesIO()
    Image.new("RGB", (1, 1), "#E6EAF0").save(stream, format="PNG")
    stream.seek(0)
    return stream


def geometry(shape) -> dict:
    return {"x_cm": round(shape.left / EMU_CM, 4), "y_cm": round(shape.top / EMU_CM, 4),
            "width_cm": round(shape.width / EMU_CM, 4), "height_cm": round(shape.height / EMU_CM, 4)}


def get_shape(slide, shape_id: int):
    found = [shape for shape in slide.shapes if shape.shape_id == shape_id]
    if len(found) != 1:
        raise PptError(f"页面中找不到唯一的顶层 shape_id={shape_id}；组合中的输入需先在本地取消组合")
    return found[0]


def shape_kind(shape) -> str | None:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape.has_table:
        return "table"
    if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
        return "image"
    if shape.has_text_frame:
        return "text"
    return None


def inner_shapes(shape):
    """All descendant shapes of a group, in document order."""
    found = []

    def walk(shapes):
        for child in shapes:
            found.append(child)
            if child.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(child.shapes)

    walk(shape.shapes)
    return found


def group_field_indexes(shape) -> list[int]:
    """Indexes (into inner_shapes) of descendants with sample text — the per-item fill fields."""
    return [i for i, child in enumerate(inner_shapes(shape))
            if child.has_text_frame and child.text_frame.text.strip()]


def effective_font(shape) -> float:
    if shape.has_table:
        frames = [shape.table.cell(0, 0).text_frame]
    elif shape.has_text_frame:
        frames = [shape.text_frame]
    else:
        return 20.0
    def explicit_size(tf):
        for p in tf.paragraphs:
            for run in p.runs:
                if run.font.size is not None:
                    return run.font.size.pt
            if p.font.size is not None:
                return p.font.size.pt

    for tf in frames:
        size = explicit_size(tf)
        if size is not None:
            return size
    # A reference often defines font size in its layout or master, not its runs.
    # Follow the same inheritance chain rather than assuming a 20pt title.
    drawing = "http://schemas.openxmlformats.org/drawingml/2006/main"
    presentation = "http://schemas.openxmlformats.org/presentationml/2006/main"
    current = shape
    level = frames[0].paragraphs[0].level + 1
    seen = set()
    while current is not None and id(current) not in seen:
        seen.add(id(current))
        if current.has_text_frame:
            size = explicit_size(current.text_frame)
            if size is not None:
                return size
            props = current.text_frame._txBody.find(f"{{{drawing}}}lstStyle/{{{drawing}}}lvl{level}pPr/{{{drawing}}}defRPr")
            if props is not None and props.get("sz"):
                return int(props.get("sz")) / 100
        current = getattr(current, "_base_placeholder", None)
    if shape.has_text_frame:
        part = shape.part
        master = part.slide_layout.slide_master
        title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
        kind = "titleStyle" if shape.is_placeholder and shape.placeholder_format.type in title_types else "bodyStyle" if shape.is_placeholder else "otherStyle"
        props = master._element.find(f"{{{presentation}}}txStyles/{{{presentation}}}{kind}/{{{drawing}}}lvl{level}pPr/{{{drawing}}}defRPr")
        if props is not None and props.get("sz"):
            return int(props.get("sz")) / 100
        default = part.package.presentation_part._element.find(f"{{{presentation}}}defaultTextStyle/{{{drawing}}}lvl{level}pPr/{{{drawing}}}defRPr")
        if default is not None and default.get("sz"):
            return int(default.get("sz")) / 100
    # Tables may rely on an application table style with no stored font size.
    return 20.0


def inferred_limits(shape, kind: str) -> dict:
    box = geometry(shape)
    base = {"max_width_cm": box["width_cm"], "max_height_cm": box["height_cm"]}
    if kind == "text":
        font = effective_font(shape)
        tf = shape.text_frame
        width_pt = max(1, (shape.width - tf.margin_left - tf.margin_right) / 12700)
        height_pt = max(1, (shape.height - tf.margin_top - tf.margin_bottom) / 12700)
        lines = max(1, int(height_pt / (font * 1.25)))
        base.update(max_chars=max(1, int(width_pt / font) * lines), max_lines=lines,
                    max_font_pt=round(font, 2))
    elif kind == "image":
        base.update(max_pixels=40_000_000)
    elif kind == "table":
        base.update(max_rows=len(shape.table.rows), max_cols=len(shape.table.columns),
                    max_cell_chars=30, max_cell_lines=2, max_font_pt=round(effective_font(shape), 2))
    return base


def inspect_reference(path: Path) -> dict:
    data = path.read_bytes()
    inspect_package(data)
    prs = Presentation(BytesIO(data))
    layouts = []
    for index, slide in enumerate(prs.slides, start=1):
        inputs = []
        for shape in slide.shapes:
            kind = shape_kind(shape)
            if kind is None:
                continue
            entry = {"id": f"element_{shape.shape_id}", "shape_id": shape.shape_id, "type": kind,
                     "description": f"{shape.name}：请根据参考页填写该输入的用途", "required": True,
                     "limits": inferred_limits(shape, kind)}
            if kind == "group":
                entry["description"] = (f"{shape.name}：组合形状；与另一个相同的组合组成 array 输入的两个样例"
                                        f"（合并为单个输入并改用 shape_ids），或删除映射保持静态")
            inputs.append(entry)
        layouts.append({"id": f"layout_{index}", "name": f"参考第 {index} 页", "source_slide": index,
                        "description": "请描述页面用途，并删除应保持静态的元素映射", "inputs": inputs})
    return {"schema_version": 1, "name": path.stem, "layouts": layouts}


def validate_spec(spec: dict, prs) -> None:
    require_keys(spec, {"schema_version", "name", "layouts"}, "模板 spec")
    if spec.get("schema_version") != 1 or not isinstance(spec.get("name"), str) or not spec["name"].strip():
        raise PptError("模板 spec 需要 schema_version: 1 和非空 name")
    layouts = spec.get("layouts")
    if not isinstance(layouts, list) or not layouts:
        raise PptError("spec.layouts 必须是非空列表")
    ids = set()
    for layout in layouts:
        if not isinstance(layout, dict):
            raise PptError("布局必须为对象")
        require_keys(layout, {"id", "name", "source_slide", "description", "inputs"}, "布局")
        lid = identifier(layout.get("id"), "布局 ID")
        if lid in ids:
            raise PptError(f"布局 ID 重复：{lid}")
        ids.add(lid)
        for key in ("name", "description"):
            if not isinstance(layout.get(key), str) or not layout[key].strip():
                raise PptError(f"{lid}.{key} 必须为非空字符串")
        idx = layout.get("source_slide")
        if isinstance(idx, bool) or not isinstance(idx, int) or not 1 <= idx <= len(prs.slides):
            raise PptError(f"{lid}.source_slide 超出参考页范围")
        slots = layout.get("inputs")
        if not isinstance(slots, list):
            raise PptError(f"{lid}.inputs 必须为列表")
        slot_ids, shape_ids = set(), set()
        for slot in slots:
            if not isinstance(slot, dict):
                raise PptError(f"{lid} 输入必须是对象")
            kind = slot.get("type")
            is_array = kind == "array"
            require_keys(slot, {"id", "shape_ids" if is_array else "shape_id", "type",
                                "description", "required", "limits"}, f"{lid} 输入")
            sid = identifier(slot.get("id"), "输入 ID")
            if sid in slot_ids:
                raise PptError(f"{lid}/{sid} 输入 ID 重复")
            slot_ids.add(sid)
            slide = prs.slides[idx - 1]
            if is_array:
                mold_ids = slot.get("shape_ids")
                if (not isinstance(mold_ids, list) or len(mold_ids) != 2 or mold_ids[0] == mold_ids[1]
                        or any(isinstance(i, bool) or not isinstance(i, int) for i in mold_ids)):
                    raise PptError(f"{lid}/{sid}.shape_ids 必须为两个不同的整数形状 ID（排列方向和间距由这两个样例的位置差决定）")
                if shape_ids & set(mold_ids):
                    raise PptError(f"{lid}/{sid} 输入 shape_id 重复")
                shape_ids.update(mold_ids)
                molds = [get_shape(slide, i) for i in mold_ids]
            else:
                shape_id = slot.get("shape_id")
                if shape_id in shape_ids:
                    raise PptError(f"{lid}/{sid} 输入 ID 或 shape_id 重复")
                if isinstance(shape_id, bool) or not isinstance(shape_id, int):
                    raise PptError(f"{lid}/{sid}.shape_id 必须为整数")
                shape_ids.add(shape_id)
                molds = [get_shape(slide, shape_id)]
            if not isinstance(slot.get("description"), str) or not slot["description"].strip():
                raise PptError(f"{lid}/{sid}.description 必须为非空字符串")
            if not isinstance(slot.get("required"), bool):
                raise PptError(f"{lid}/{sid}.required 必须为布尔值")
            limits = slot.get("limits")
            if not isinstance(limits, dict):
                raise PptError(f"{lid}/{sid} 缺少 limits")
            if is_array:
                validate_array(lid, sid, slot, molds, prs)
                continue
            shape = molds[0]
            if kind not in {"text", "image", "table"} or kind != shape_kind(shape):
                raise PptError(f"{lid}/{sid} 输入类型与实际形状不匹配")
            necessary = {"max_width_cm", "max_height_cm"}
            integers = set()
            if kind == "text":
                necessary |= {"max_chars", "max_lines", "max_font_pt"}
                integers = {"max_chars", "max_lines"}
            elif kind == "image":
                necessary |= {"max_pixels"}
                integers = {"max_pixels"}
            else:
                necessary |= {"max_rows", "max_cols", "max_cell_chars", "max_cell_lines", "max_font_pt"}
                integers = {"max_rows", "max_cols", "max_cell_chars", "max_cell_lines"}
            require_keys(limits, necessary, f"{lid}/{sid}.limits")
            if necessary - set(limits):
                raise PptError(f"{lid}/{sid} 缺少约束：{sorted(necessary - set(limits))}")
            for key in necessary:
                positive(limits[key], f"{lid}/{sid}/{key}", integer=key in integers)
            box = geometry(shape)
            if limits["max_width_cm"] > box["width_cm"] + .001 or limits["max_height_cm"] > box["height_cm"] + .001:
                raise PptError(f"{lid}/{sid} 尺寸约束不得超过参考输入框")
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + 360 or shape.top + shape.height > prs.slide_height + 360:
                raise PptError(f"{lid}/{sid} 输入框位于页面外")
            if kind == "table":
                if limits["max_rows"] > len(shape.table.rows) or limits["max_cols"] > len(shape.table.columns):
                    raise PptError(f"{lid}/{sid} 表格容量超过参考表格网格；请先在参考中扩大网格")
                if any(shape.table.cell(r, c).is_merge_origin or shape.table.cell(r, c).is_spanned
                       for r in range(len(shape.table.rows)) for c in range(len(shape.table.columns))):
                    raise PptError(f"{lid}/{sid} 输入表格不能含合并单元格")


def validate_array(lid: str, sid: str, slot: dict, molds: list, prs) -> None:
    """array 输入：两个样例形状决定条目类型与排列向量（方向+间距）。"""
    kinds = {shape_kind(m) for m in molds}
    if len(kinds) != 1 or not kinds <= {"text", "image", "group"}:
        raise PptError(f"{lid}/{sid} 数组样例类型须一致，且为文本框、图片或组合")
    item_kind = kinds.pop()
    if abs(molds[0].width - molds[1].width) > 360 or abs(molds[0].height - molds[1].height) > 360:
        raise PptError(f"{lid}/{sid} 数组两个样例大小必须一致")
    step = (molds[1].left - molds[0].left, molds[1].top - molds[0].top)
    if abs(step[0]) + abs(step[1]) <= 360:
        raise PptError(f"{lid}/{sid} 数组两个样例必须位于不同位置，其位置差定义排列方向和间距")
    if item_kind == "group":
        counts = [len([c for c in inner_shapes(m) if c.has_text_frame]) for m in molds]
        if counts[0] != counts[1]:
            raise PptError(f"{lid}/{sid} 两个组合样例的内部文本形状数量必须一致")
    limits = slot["limits"]
    necessary = {"max_items"}
    integers = {"min_items", "max_items"}
    optional = {"min_items"}
    if item_kind == "image":
        necessary |= {"max_item_pixels"}
        integers.add("max_item_pixels")
    else:
        necessary |= {"max_item_chars", "max_item_lines", "max_item_font_pt"}
        integers |= {"max_item_chars", "max_item_lines"}
    require_keys(limits, necessary | optional, f"{lid}/{sid}.limits")
    if necessary - set(limits):
        raise PptError(f"{lid}/{sid} 缺少约束：{sorted(necessary - set(limits))}")
    for key in necessary:
        positive(limits[key], f"{lid}/{sid}/{key}", integer=key in integers)
    minimum = limits.get("min_items", 1)
    if "min_items" in limits:
        positive(minimum, f"{lid}/{sid}/min_items", integer=True)
        if minimum > limits["max_items"]:
            raise PptError(f"{lid}/{sid} min_items 不得大于 max_items")
    # 条目位置按两个样例的间距向量外推；最大个数必须整体落在页面内。
    for i in range(int(limits["max_items"])):
        left = molds[0].left + step[0] * i
        top = molds[0].top + step[1] * i
        if left < -360 or top < -360 or left + molds[0].width > prs.slide_width + 360 or top + molds[0].height > prs.slide_height + 360:
            raise PptError(f"{lid}/{sid} 按间距排列第 {i + 1} 个条目将超出页面，请缩小间距或降低 max_items")


def clear_frame(tf) -> None:
    """Keep first paragraph/run defaults, remove sample text and hyperlinks."""
    prototype = tf.paragraphs[0]
    run_props = deepcopy(prototype.runs[0]._r.get_or_add_rPr()) if prototype.runs else None
    tf.clear()
    paragraph = tf.paragraphs[0]
    run = paragraph.add_run()
    if run_props is not None:
        for node in list(run_props):
            if node.tag.endswith(("}hlinkClick", "}hlinkMouseOver")):
                run_props.remove(node)
        run._r.insert(0, run_props)
    run.text = ""


def remove_sample(shape, slide) -> None:
    kind = shape_kind(shape)
    if kind == "group":
        for child in inner_shapes(shape):
            if child.has_text_frame:
                clear_frame(child.text_frame)
    elif kind == "text":
        clear_frame(shape.text_frame)
    elif kind == "table":
        for row in shape.table.rows:
            for cell in row.cells:
                clear_frame(cell.text_frame)
    else:
        _, rid = slide.part.get_or_add_image_part(blank_image())
        blip = shape._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
        blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)
        shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
        # Trim the old unused relationship so source sample images aren't retained.
        used = set(shape_id for node in slide._element.iter() for attr, shape_id in node.attrib.items()
                   if attr.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"))
        for rel in list(slide.part.rels.values()):
            if rel.reltype.endswith("/image") and rel.rId not in used:
                slide.part.drop_rel(rel.rId)


def documentation(manifest: dict) -> str:
    lines = [f"# {manifest['name']} — PPT 模板说明", "",
             f"画布：{manifest['slide_size_cm']['width']:.2f} × {manifest['slide_size_cm']['height']:.2f} cm", "",
             "本说明由模板清单自动生成。每个布局是一张示例页面，可重复使用并任意排序。",
             "字符数按解析样式后的 Unicode 字符计算，含空格和标点；样式标记不计入。",
             "表格行数包含表头。尺寸单位为 cm，字号单位为 pt。",
             "字数、估算行数和尺寸是生成前约束；实际字体换行仍须使用本地 PowerPoint/LibreOffice 复核。", ""]
    for layout in manifest["layouts"]:
        lines += [f"## {layout['id']}：{layout['name']}", "", layout["description"], "",
                  f"模板示例页：{layout['template_slide']}。", ""]
        for slot in layout["inputs"]:
            box, limits = slot["box_cm"], slot["limits"]
            if slot["type"] == "array":
                step = slot["step_cm"]
                direction = "横向" if abs(step["x_cm"]) >= abs(step["y_cm"]) else "纵向"
                item_desc = {"text": "单个文本形状", "image": "本地图片", "group": "组合形状"}[slot["item_kind"]]
                fields = {"text": 1, "image": 0}.get(slot["item_kind"],
                                                     len(slot.get("field_indexes") or []))
                lines += [f"### {slot['id']}（array，{'建议填写' if slot['required'] else '可省略'}）", "",
                          slot["description"], "",
                          f"映射：shape_ids={slot['shape_ids']}（两个样例定义{direction}排列，"
                          f"间距 Δx={step['x_cm']} cm，Δy={step['y_cm']} cm）；"
                          f"首条目位置 x={box['x_cm']:.2f}，y={box['y_cm']:.2f}；"
                          f"条目框 {box['width_cm']:.2f} × {box['height_cm']:.2f} cm。", "",
                          "填充时按条目个数沿样例间距复制排列；未提供时移除全部样例，不会阻止生成。", "",
                          f"- 条目类型：{item_desc}；条目数 {limits.get('min_items', 1)}..{limits['max_items']} 个。"]
                if slot["item_kind"] == "image":
                    lines += ["- 每条目一张本地图片：`- 图片路径` 或 `- {\"src\":\"路径\",\"fit\":\"cover\"}`；"
                              f"最多 {limits['max_item_pixels']} 像素；条目框内按 contain/cover 适配。"]
                else:
                    lines += [f"- 每条目 {fields} 个文本位：`- 文本`（单文本位）或 `- 文本位1 || 文本位2`；"
                              f"每个文本位最多 {limits['max_item_chars']} 字符、{limits['max_item_lines']} 行，"
                              f"片段最大字号 {limits['max_item_font_pt']} pt，支持局部文字样式。",
                              "- 也可用 JSON：```{\"items\":[[\"字段1\",\"字段2\"], ...]}```"]
                lines.append("")
                continue
            lines += [f"### {slot['id']}（{slot['type']}，{'建议填写' if slot['required'] else '可省略'}）", "",
                      slot["description"], "",
                      f"映射：shape_id={slot['shape_id']}；位置 x={box['x_cm']:.2f}，y={box['y_cm']:.2f}；"
                      f"输入框 {box['width_cm']:.2f} × {box['height_cm']:.2f} cm。", "",
                      "未在填充内容中提供时按空内容处理，不会阻止生成。", "",
                      f"- 最大宽度：{limits['max_width_cm']} cm；最大高度：{limits['max_height_cm']} cm。"]
            if slot["type"] == "text":
                lines += [f"- 最多 {limits['max_chars']} 字符；最多 {limits['max_lines']} 行（含自动换行估算）；"
                          f"片段最大字号 {limits['max_font_pt']} pt。",
                          "- 输入普通文字、**加粗**、*斜体*、span 局部字体/字号/颜色/下划线。未覆盖的样式继承模板。"]
            elif slot["type"] == "image":
                lines += [f"- 本地 PNG/JPEG 等 Pillow 支持的位图，最多 {limits['max_pixels']} 像素。",
                          "- contain 保持完整图片；cover 居中裁剪。可在上限内指定宽高，图片居中于输入框。"]
            else:
                lines += [f"- 最多 {limits['max_rows']} 行 × {limits['max_cols']} 列；单元格最多 {limits['max_cell_chars']} 字符、"
                          f"{limits['max_cell_lines']} 行；片段最大字号 {limits['max_font_pt']} pt。",
                          "- 输入 Markdown 表格或 JSON rows；每个单元格支持局部文字样式。未用网格行列会收起。"]
            lines.append("")
    lines += ["## 本地处理边界", "", "生成脚本没有网络请求、上传、遥测或远程资源下载。",
              "拒绝 URL、UNC 共享、PPTX 外部关系、宏、ActiveX 和输入表格合并单元格。",
              "静态内容（未列为输入的形状、母版、主题、背景、图表等）会保留，发布模板前请检查其中是否含原稿信息。",
              "模板清单与 PPTX 通过 SHA-256 绑定。修改 PPTX 后应重新创建模板；只修改约束时重新执行 describe。", ""]
    return "\n".join(lines)


def create_template(reference: Path, spec: dict, output: Path) -> dict:
    if output.exists():
        raise PptError(f"模板目录已存在，请使用新的目录：{output}")
    data = reference.read_bytes()
    inspect_package(data)
    source = Presentation(BytesIO(data))
    validate_spec(spec, source)
    prs = Presentation(BytesIO(clone_slides(data, [l["source_slide"] for l in spec["layouts"]])))
    manifest = {"schema_version": 1, "name": spec["name"], "template_file": "template.pptx",
                "slide_size_cm": {"width": round(prs.slide_width / EMU_CM, 4), "height": round(prs.slide_height / EMU_CM, 4)},
                "layouts": []}
    for index, layout in enumerate(spec["layouts"], start=1):
        entry = deepcopy(layout)
        entry.pop("source_slide")
        entry["template_slide"] = index
        for slot in entry["inputs"]:
            if slot["type"] == "array":
                molds = [get_shape(prs.slides[index - 1], i) for i in slot["shape_ids"]]
                slot["item_kind"] = shape_kind(molds[0])
                slot["box_cm"] = geometry(molds[0])
                slot["step_cm"] = {"x_cm": round((molds[1].left - molds[0].left) / EMU_CM, 4),
                                   "y_cm": round((molds[1].top - molds[0].top) / EMU_CM, 4)}
                if slot["item_kind"] == "group":
                    fields = [group_field_indexes(m) for m in molds]
                    if not fields[0] or len(fields[0]) != len(fields[1]):
                        raise PptError(f"{layout['id']}/{slot['id']} 组合样例须含相同数量的非空文本形状（空的装饰形状会成为条目静态部分）")
                    slot["field_indexes"] = fields[0]
                slot["default_font_pt"] = effective_font(molds[0])
                for mold in molds:
                    remove_sample(mold, prs.slides[index - 1])
            else:
                shape = get_shape(prs.slides[index - 1], slot["shape_id"])
                slot["box_cm"] = geometry(shape)
                slot["default_font_pt"] = effective_font(shape)
                remove_sample(shape, prs.slides[index - 1])
        manifest["layouts"].append(entry)
    # Reduce metadata carryover; static objects are deliberately retained.
    for attr in ("title", "subject", "author", "keywords", "comments", "last_modified_by", "category"):
        setattr(prs.core_properties, attr, "")
    buffer = BytesIO()
    prs.save(buffer)
    pptx = buffer.getvalue()
    inspect_package(pptx)
    manifest["template_sha256"] = sha256(pptx)
    output.parent.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=".ppt-template-", dir=output.parent))
    try:
        (staging / "template.pptx").write_bytes(pptx)
        (staging / "manifest.json").write_bytes(json_bytes(manifest))
        (staging / "template-guide.md").write_text(documentation(manifest), encoding="utf-8")
        staging.rename(output)
    finally:
        if staging.exists():
            # Verify the resolved target before recursively removing a failed staging directory.
            if staging.resolve().parent != output.parent.resolve() or not staging.name.startswith(".ppt-template-"):
                raise PptError("暂存目录超出预期位置，停止清理")
            shutil.rmtree(staging)
    return manifest


def load_template(directory: Path):
    manifest = read_json(directory / "manifest.json")
    if manifest.get("template_file") != "template.pptx":
        raise PptError("manifest.template_file 必须为 template.pptx")
    data = (directory / "template.pptx").read_bytes()
    if sha256(data) != manifest.get("template_sha256"):
        raise PptError("模板 PPTX 与 manifest SHA-256 不一致，请重新创建模板")
    inspect_package(data)
    prs = Presentation(BytesIO(data))
    expected_size = {"width": round(prs.slide_width / EMU_CM, 4), "height": round(prs.slide_height / EMU_CM, 4)}
    if manifest.get("slide_size_cm") != expected_size:
        raise PptError("manifest 画布尺寸与 PPTX 不一致")
    # Reuse the strict spec validator after stripping derived manifest fields.
    spec = {"schema_version": manifest.get("schema_version"), "name": manifest.get("name"), "layouts": []}
    for layout in manifest.get("layouts", []):
        entry = {key: layout.get(key) for key in ("id", "name", "description", "inputs")}
        entry["source_slide"] = layout.get("template_slide")
        entry["inputs"] = [{k: s.get(k) for k in
                            ("id", "shape_ids" if s.get("type") == "array" else "shape_id",
                             "type", "description", "required", "limits")}
                           for s in entry["inputs"]]
        spec["layouts"].append(entry)
    validate_spec(spec, prs)
    for layout in manifest["layouts"]:
        slide = prs.slides[layout["template_slide"] - 1]
        for slot in layout["inputs"]:
            if slot["type"] == "array":
                molds = [get_shape(slide, i) for i in slot["shape_ids"]]
                if slot.get("item_kind") != shape_kind(molds[0]):
                    raise PptError(f"{layout['id']}/{slot['id']} 数组条目类型与 PPTX 不一致")
                if slot.get("box_cm") != geometry(molds[0]):
                    raise PptError(f"{layout['id']}/{slot['id']} 输入几何信息与 PPTX 不一致")
                step = slot.get("step_cm", {})
                expected = {"x_cm": round(slot["box_cm"]["x_cm"] + step.get("x_cm", 0), 4),
                            "y_cm": round(slot["box_cm"]["y_cm"] + step.get("y_cm", 0), 4),
                            "width_cm": slot["box_cm"]["width_cm"], "height_cm": slot["box_cm"]["height_cm"]}
                if geometry(molds[1]) != expected:
                    raise PptError(f"{layout['id']}/{slot['id']} 第二个数组样例位置与记录间距不一致")
                if slot["item_kind"] == "group":
                    fields = slot.get("field_indexes")
                    descendants = inner_shapes(molds[0])
                    if (not isinstance(fields, list) or not fields
                            or any(not isinstance(i, int) or not 0 <= i < len(descendants)
                                   or not descendants[i].has_text_frame for i in fields)):
                        raise PptError(f"{layout['id']}/{slot['id']} field_indexes 与 PPTX 组合结构不一致")
                if slot.get("default_font_pt") != effective_font(molds[0]):
                    raise PptError(f"{layout['id']}/{slot['id']} 默认字号与 PPTX 不一致")
                continue
            shape = get_shape(slide, slot["shape_id"])
            if slot.get("box_cm") != geometry(shape):
                raise PptError(f"{layout['id']}/{slot['id']} 输入几何信息与 PPTX 不一致")
            if slot.get("default_font_pt") != effective_font(shape):
                raise PptError(f"{layout['id']}/{slot['id']} 默认字号与 PPTX 不一致")
    return manifest, data


def starter_reference() -> tuple[bytes, dict]:
    """内置模板：浅色简约风格（白底 + 橙红强调 + 浅灰卡片），18 种布局示例。

    数组输入由两个样例形状定义：它们的位置差决定条目的排列方向和间距，
    填充内容按条目个数沿该向量复制排列。
    """
    INK, ORANGE, MUTED, CARD, LINE, WHITE = "22242E", "E85A2E", "8A8D98", "F5F4F1", "E8E4DE", "FFFFFF"
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    layouts = []

    def block(shapes, x, y, w, h, fill, *, round_=False, radius=0.12, text_sample=None,
              size=14, color=WHITE, bold=False):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else (MSO_SHAPE.OVAL if fill == "oval" else MSO_SHAPE.RECTANGLE)
        shape = shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill if fill != "oval" else ORANGE)
        shape.line.fill.background()
        shape.shadow.inherit = False
        if round_:
            shape.adjustments[0] = radius
        tf = shape.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = tf.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = text_sample or ""
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        return shape

    def text(shapes, name, x, y, w, h, size, color=INK, *, bold=False, sample=None):
        shape = shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
        shape.name = name
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Cm(.1)
        tf.margin_top = tf.margin_bottom = Cm(.05)
        run = tf.paragraphs[0].add_run()
        run.text = sample if sample is not None else name
        run.font.name, run.font.size, run.font.bold = "Microsoft YaHei", Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
        return shape

    def deco(slide):
        """内容页通用装饰：标题上方的小橙条 + 右上角标块 + 底部细线。"""
        block(slide.shapes, 2.05, 0.78, 0.9, 0.13, ORANGE, round_=True, radius=0.5)
        block(slide.shapes, 31.6, 0.9, 0.85, 0.85, ORANGE, round_=True, radius=0.3)
        block(slide.shapes, 2.0, 17.9, 29.87, 0.03, LINE)

    def tslot(sid, shape, desc, required, chars, lines, font, w, h):
        return {"id": sid, "shape_id": shape.shape_id, "type": "text", "description": desc,
                "required": required, "limits": {"max_chars": chars, "max_lines": lines, "max_font_pt": font,
                                                 "max_width_cm": w, "max_height_cm": h}}

    def islot(sid, shape, desc, required, w, h):
        return {"id": sid, "shape_id": shape.shape_id, "type": "image", "description": desc,
                "required": required, "limits": {"max_width_cm": w, "max_height_cm": h, "max_pixels": 40_000_000}}

    def aslot(sid, molds, desc, required, max_items, *, min_items=1, chars=None, lines=None,
              font=None, pixels=None):
        limits = {"min_items": min_items, "max_items": max_items}
        if pixels is not None:
            limits["max_item_pixels"] = pixels
        else:
            limits.update(max_item_chars=chars, max_item_lines=lines, max_item_font_pt=font)
        return {"id": sid, "shape_ids": [m.shape_id for m in molds], "type": "array",
                "description": desc, "required": required, "limits": limits}

    def header(slide, slots, *, chars=24):
        slots.append(tslot("title", text(slide.shapes, "title", 2.0, 0.95, 22, 1.5, 26, bold=True),
                           "页面标题", True, chars, 1, 26, 22, 1.5))
        slots.append(tslot("subtitle", text(slide.shapes, "subtitle", 2.0, 2.5, 24, 0.9, 13, MUTED),
                           "副标题/导读，可省略", False, 42, 1, 13, 24, 0.9))
        deco(slide)

    def pic(slide, x, y, w, h):
        shape = slide.shapes.add_picture(blank_image(), Cm(x), Cm(y), Cm(w), Cm(h))
        shape.name = "image"
        return shape

    def band(slide):
        block(slide.shapes, 0, 16.0, 33.867, 3.05, ORANGE)
        block(slide.shapes, 8.8, 17.05, 2.8, 0.95, WHITE, round_=True, radius=0.5)

    def logo(slide):
        block(slide.shapes, 2.2, 1.5, 3.4, 1.05, CARD, round_=True, radius=0.5)
        block(slide.shapes, 2.5, 1.78, 0.5, 0.5, "oval")

    def group_of(slide, builders):
        shape = slide.shapes.add_group_shape()
        for build in builders:
            build(shape.shapes)
        return shape

    def chip(x, y, d, sample, *, size=15):
        return lambda s: block(s, x, y, d, d, ORANGE, round_=True, radius=0.35,
                               text_sample=sample, size=size, bold=True)

    def card_parts(x, y, w, h, samples):
        bg = lambda s: block(s, x, y, w, h, CARD, round_=True)
        return [bg] + samples(x, y, w, h)

    slide_of = lambda: prs.slides.add_slide(prs.slide_layouts[6])
    index = 0

    def emit(lid, name, desc, slide, slots):
        nonlocal index
        index += 1
        layouts.append({"id": lid, "name": name, "source_slide": index, "description": desc, "inputs": slots})

    # 1 封面：左文右图 + 底部橙色横带
    slide = slide_of()
    band(slide)
    logo(slide)
    emit("cover", "封面页", "演示开场：大标题 + 副标题 + 底部横带上的日期/署名，右侧可配图", slide, [
        tslot("title", text(slide.shapes, "title", 2.2, 6.0, 17.6, 4.3, 40, bold=True),
              "演示主题，最多两行", True, 24, 2, 40, 17.6, 4.3),
        tslot("subtitle", text(slide.shapes, "subtitle", 2.2, 10.9, 17, 1.1, 17, MUTED),
              "一句话副标题，可省略", False, 32, 1, 17, 17, 1.1),
        tslot("meta", text(slide.shapes, "meta", 2.2, 17.05, 15, 0.95, 12, WHITE),
              "横带上的日期/署名，可省略", False, 32, 1, 12, 15, 0.95),
        islot("image", pic(slide, 20.4, 2.3, 11.4, 11.0), "右侧配图，可省略", False, 11.4, 11.0)])

    # 2 目录：左图右列表，条目为「序号角标 + 名称」组合数组
    slide = slide_of()
    molds = [
        group_of(slide, [chip(15.0, 5.8 + dy, 1.15, str(n)),
                         lambda s, dy=dy, n=n: text(s, "item", 16.6, 5.9 + dy, 14.6, 0.95, 17, INK,
                                                  sample=f"目录项{n}")])
        for n, dy in ((1, 0.0), (2, 2.0))]
    emit("agenda", "目录页", "左侧配图与说明，右侧编号目录项（数组条目：序号 || 名称，纵向排列，最多 6 项）", slide, [
        islot("image", pic(slide, 2.0, 3.6, 10.4, 10.8), "左侧配图，可省略", False, 10.4, 10.8),
        tslot("caption", text(slide.shapes, "caption", 2.0, 14.8, 10.4, 1.1, 15, INK, bold=True, sample="Agenda."),
              "图片下方说明（如 Agenda.），可省略", False, 14, 1, 15, 10.4, 1.1),
        tslot("title", text(slide.shapes, "title", 15.0, 3.6, 16, 1.6, 26, bold=True, sample="目录"),
              "栏题，如「目录」", True, 14, 1, 26, 16, 1.6),
        aslot("items", molds, "目录条目：- 序号 || 名称", True, 6, chars=24, lines=1, font=17)])

    # 3 章节页：胶囊标签 + 大标题 + 说明，右上角色块与可选配图
    slide = slide_of()
    block(slide.shapes, 27.0, 0, 6.867, 4.4, ORANGE)
    tag = block(slide.shapes, 2.2, 6.4, 4.8, 1.15, ORANGE, round_=True, radius=0.5,
                text_sample="PART 01", size=14, bold=True)
    emit("section", "章节过渡页", "章节分隔：胶囊标签 + 章节标题 + 说明，右上角可选配图", slide, [
        tslot("tag", tag, "胶囊标签，如 PART 01，可省略", False, 12, 1, 14, 4.8, 1.15),
        tslot("title", text(slide.shapes, "title", 2.2, 8.1, 19.5, 2.4, 34, bold=True),
              "章节标题", True, 16, 1, 34, 19.5, 2.4),
        tslot("desc", text(slide.shapes, "desc", 2.2, 10.9, 18, 2.2, 15, MUTED),
              "章节导读，可省略", False, 70, 3, 15, 18, 2.2),
        islot("image", pic(slide, 21.5, 1.7, 9.9, 6.2), "右上角配图，可省略", False, 9.9, 6.2)])

    # 4 标题正文
    slide = slide_of()
    slots = []
    header(slide, slots)
    slots.append(tslot("body", text(slide.shapes, "body", 2.0, 3.9, 29.8, 13.3, 17),
                       "正文：普通段落、- 符号分点、1. 数字分点、- [图标] 图标分点", True, 430, 13, 17, 29.8, 13.3))
    emit("title_body", "标题正文页", "标题 + 副标题 + 大段正文，支持三种分点与局部样式", slide, slots)

    # 5 双栏卡片
    slide = slide_of()
    slots = []
    header(slide, slots)
    for x in (2.0, 17.52):
        block(slide.shapes, x, 4.0, 14.35, 12.9, CARD, round_=True)
        block(slide.shapes, x + 0.55, 4.55, 1.3, 0.4, ORANGE, round_=True, radius=0.5)
    slots.append(tslot("left", text(slide.shapes, "left", 2.7, 5.4, 13.0, 10.9, 15),
                       "左栏内容（可用分点）", True, 260, 11, 15, 13.0, 10.9))
    slots.append(tslot("right", text(slide.shapes, "right", 18.2, 5.4, 13.0, 10.9, 15),
                       "右栏内容（可用分点）", True, 260, 11, 15, 13.0, 10.9))
    emit("two_column", "双栏卡片页", "标题 + 左右两块浅灰卡片文字，适合并列的两组内容", slide, slots)

    # 6 左图右文
    slide = slide_of()
    slots = []
    header(slide, slots)
    slots.append(islot("image", pic(slide, 2.0, 4.0, 13.2, 12.9), "左侧配图", True, 13.2, 12.9))
    slots.append(tslot("body", text(slide.shapes, "body", 17.0, 4.0, 14.8, 12.9, 15),
                       "右侧文字（可用分点）", True, 300, 12, 15, 14.8, 12.9))
    emit("image_text", "左图右文页", "标题 + 左侧图片 + 右侧说明文字", slide, slots)

    # 7 左文右图
    slide = slide_of()
    slots = []
    header(slide, slots)
    slots.append(tslot("body", text(slide.shapes, "body", 2.0, 4.0, 14.8, 12.9, 15),
                       "左侧文字（可用分点）", True, 300, 12, 15, 14.8, 12.9))
    slots.append(islot("image", pic(slide, 18.2, 4.0, 13.2, 12.9), "右侧配图", True, 13.2, 12.9))
    emit("text_image", "左文右图页", "标题 + 左侧说明文字 + 右侧图片", slide, slots)

    # 8 卡片阵列：两行数组，每条目为「图标角标 + 标题 + 说明」卡片
    slide = slide_of()
    slots = []
    header(slide, slots)

    def card_row(x0, y):
        return [
            group_of(slide, card_parts(x, y, 7.05, 5.7, lambda x, y, w, h, n=n: [
                lambda s: block(s, x + 0.55, y + 0.6, 1.25, 1.25, ORANGE, round_=True, radius=0.35,
                                text_sample="✓", size=15, bold=True),
                lambda s: text(s, "card_title", x + 0.55, y + 2.3, 5.9, 0.9, 15, INK, bold=True,
                               sample="卡片标题"),
                lambda s: text(s, "card_desc", x + 0.55, y + 3.25, 5.95, 1.9, 11, MUTED,
                               sample="卡片说明文字")]))
            for n, x in enumerate((x0, x0 + 7.55), start=1)]
    emit("cards", "卡片阵列页", "标题 + 至多两排特性卡片（数组条目：图标 || 标题 || 说明，横向排列，每排最多 4 张；第二排可省略）",
         slide, slots + [
             aslot("row_a", card_row(2.0, 4.2), "第一排卡片：- 图标 || 标题 || 说明", True, 4,
                   chars=40, lines=2, font=15),
             aslot("row_b", card_row(2.0, 10.5), "第二排卡片，可省略", False, 4,
                   chars=40, lines=2, font=15)])

    # 9 步骤条：编号角标 + 阶段 + 说明，横向数组
    slide = slide_of()
    slots = []
    header(slide, slots)
    molds = [
        group_of(slide, [chip(x, 4.6, 1.25, str(n)),
                         lambda s, x=x, n=n: text(s, "stage", x, 6.25, 5.3, 0.95, 15, INK, bold=True,
                                                  sample=f"阶段{n}"),
                         lambda s, x=x, n=n: text(s, "stage_desc", x, 7.3, 5.3, 4.0, 11, MUTED,
                                                  sample="阶段说明")])
        for n, x in enumerate((2.0, 8.3), start=1)]
    slots.append(aslot("items", molds, "流程步骤：- 序号 || 阶段 || 说明（横向，最多 5 步）", True, 5,
                       chars=60, lines=4, font=15))
    emit("steps", "步骤流程页", "标题 + 编号步骤条，适合流程、阶段和计划", slide, slots)

    # 10 数据亮点：大数字 + 标签 + 说明，横向数组
    slide = slide_of()
    slots = []
    header(slide, slots)
    molds = [
        group_of(slide, [lambda s, x=x, n=n: text(s, "num", x, 5.6, 7.3, 2.4, 34, ORANGE, bold=True,
                                                  sample="30%"),
                         lambda s, x=x, n=n: text(s, "label", x, 8.2, 7.3, 0.95, 15, INK, bold=True,
                                                  sample="指标名称"),
                         lambda s, x=x, n=n: text(s, "num_desc", x, 9.3, 7.3, 2.4, 11, MUTED,
                                                  sample="指标说明")])
        for n, x in enumerate((2.0, 9.95), start=1)]
    slots.append(aslot("items", molds, "数据卡片：- 数值 || 名称 || 说明（横向，最多 4 项）", True, 4,
                       chars=40, lines=2, font=34))
    emit("stats", "数据亮点页", "标题 + 大数字指标卡，适合关键成果与量化目标", slide, slots)

    # 11 多栏对比：栏目胶囊 + 说明卡片，横向数组
    slide = slide_of()
    slots = []
    header(slide, slots)
    molds = [
        group_of(slide, [
            lambda s, x=x, n=n: block(s, x, 4.2, 5.9, 9.6, CARD, round_=True),
            lambda s, x=x, n=n: block(s, x + 0.45, 4.75, 5.0, 1.15, ORANGE, round_=True, radius=0.5,
                                      text_sample=f"栏目{n}", size=14, bold=True),
            lambda s, x=x, n=n: text(s, "col_desc", x + 0.45, 6.35, 5.0, 6.9, 12, INK,
                                     sample="栏目说明")])
        for n, x in enumerate((2.0, 8.55), start=1)]
    slots.append(aslot("items", molds, "对比栏目：- 栏目标题 || 说明（横向，最多 4 栏）", True, 4,
                       chars=90, lines=6, font=14))
    emit("columns", "多栏对比页", "标题 + 胶囊标题栏目卡，适合多维对比与分类说明", slide, slots)

    # 12 表格页：橙色表头 + 隔行浅灰
    slide = slide_of()
    slots = []
    header(slide, slots)
    shape = slide.shapes.add_table(6, 4, Cm(2.0), Cm(4.0), Cm(29.87), Cm(12.6))
    shape.name = "data"
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    for r in range(6):
        for c in range(4):
            cell = table.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_top = cell.margin_bottom = Cm(.1)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(ORANGE if r == 0 else CARD if r % 2 else WHITE)
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = "表头" if r == 0 else "单元格"
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(14 if r == 0 else 13)
            run.font.bold = r == 0
            run.font.color.rgb = RGBColor.from_string(WHITE if r == 0 else INK)
    slots.append({"id": "data", "shape_id": shape.shape_id, "type": "table",
                  "description": "第一行为表头，最多六行四列", "required": True,
                  "limits": {"max_width_cm": 29.87, "max_height_cm": 12.6, "max_rows": 6,
                             "max_cols": 4, "max_cell_chars": 30, "max_cell_lines": 2, "max_font_pt": 14}})
    emit("table", "数据表格页", "标题 + 原生可编辑表格（橙色表头、隔行浅灰）", slide, slots)

    # 13 清单页：图标角标 + 事项，纵向数组
    slide = slide_of()
    slots = []
    header(slide, slots)
    molds = [
        group_of(slide, [chip(2.0, y, 1.05, "✓", size=14),
                         lambda s, y=y, n=n: text(s, "todo", 3.55, y + 0.08, 26, 0.9, 15, INK,
                                                  sample=f"事项{n}")])
        for n, y in enumerate((4.3, 6.05), start=1)]
    slots.append(aslot("items", molds, "清单事项：- 图标 || 事项文字（纵向，最多 7 项）", True, 7,
                       chars=36, lines=1, font=15))
    emit("checklist", "要点清单页", "标题 + 纵向图标事项列表，适合结论、待办与要点汇总", slide, slots)

    # 14 图片墙：图片数组横向排列
    slide = slide_of()
    slots = []
    header(slide, slots)
    molds = [pic(slide, x, 4.6, 5.6, 9.4) for x in (2.0, 7.95)]
    for m in molds:
        m.name = "photo"
    slots.append(aslot("items", molds, "本地图片条目：- 图片路径 或 - {\"src\":...}（横向，最多 5 张）", True, 5,
                       pixels=40_000_000))
    emit("image_row", "图片阵列页", "标题 + 一排本地图片，适合截图墙、成员/Logo 展示", slide, slots)

    # 15 标签行：单文本形状（胶囊）数组
    slide = slide_of()
    slots = []
    header(slide, slots)
    molds = []
    for n, x in enumerate((2.0, 6.7), start=1):
        pill = block(slide.shapes, x, 5.0, 4.2, 1.2, CARD, round_=True, radius=0.5,
                     text_sample=f"标签{n}", size=14, color=INK)
        molds.append(pill)
    slots.append(aslot("items", molds, "标签条目：- 标签文字（横向，最多 6 个）", True, 6,
                       chars=10, lines=1, font=14))
    emit("tags", "标签行页", "标题 + 一排胶囊标签，适合关键词、技术栈与特性罗列", slide, slots)

    # 16 引用页
    slide = slide_of()
    text(slide.shapes, "quote_mark", 2.0, 4.0, 5, 5, 110, LINE, bold=True, sample="“")
    deco(slide)
    emit("quote", "引用页", "大字引述 + 出处，适合观点、用户证言与总结语", slide, [
        tslot("quote", text(slide.shapes, "quote", 5.5, 5.6, 24.5, 4.6, 24),
              "引述内容", True, 96, 3, 24, 24.5, 4.6),
        tslot("source", text(slide.shapes, "source", 5.5, 11.0, 24.5, 1.2, 14, MUTED),
              "出处/署名，可省略", False, 34, 1, 14, 24.5, 1.2)])

    # 17 全幅大图
    slide = slide_of()
    pic_shape = pic(slide, 1.0, 1.0, 31.87, 17.05)
    block(slide.shapes, 0, 15.7, 33.867, 3.35, ORANGE)
    emit("image_full", "全幅图片页", "几乎铺满的本地大图 + 底部横带说明文字", slide, [
        islot("image", pic_shape, "全幅配图（建议 cover 裁剪）", True, 31.87, 17.05),
        tslot("caption", text(slide.shapes, "caption", 2.0, 16.35, 24, 1.3, 15, WHITE),
              "横带说明文字，可省略", False, 40, 1, 15, 24, 1.3)])

    # 18 结尾页
    slide = slide_of()
    band(slide)
    logo(slide)
    emit("closing", "结尾页", "致谢语 + 联系方式行，底部横带与封面呼应", slide, [
        tslot("title", text(slide.shapes, "title", 2.2, 7.0, 17.6, 2.6, 36, bold=True, sample="感谢观看"),
              "结束语，如「感谢观看」", True, 18, 1, 36, 17.6, 2.6),
        tslot("subtitle", text(slide.shapes, "subtitle", 2.2, 10.0, 17, 1.2, 16, MUTED),
              "补充一句，可省略", False, 32, 1, 16, 17, 1.2),
        tslot("meta", text(slide.shapes, "meta", 2.2, 17.05, 15, 0.95, 12, WHITE),
              "横带上的日期/署名，可省略", False, 32, 1, 12, 15, 0.95),
        islot("image", pic(slide, 20.4, 2.3, 11.4, 11.0), "右侧配图，可省略", False, 11.4, 11.0)])

    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue(), {"schema_version": 1, "name": "本地简约演示模板", "layouts": layouts}
