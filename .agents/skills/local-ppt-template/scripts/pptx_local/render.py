from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
import math
from pathlib import Path
import unicodedata
import warnings

from PIL import Image, ImageOps, UnidentifiedImageError
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Cm, Pt

from .common import PptError, atomic_write, local_path
from .markdown import (array_content, image_content, image_value, paragraphs,
                       parse_pages, table_content)
from .package import clone_slides, inspect_package
from .template import EMU_CM, effective_font, get_shape, inner_shapes, load_template


@dataclass
class Prepared:
    page_id: str
    layout: dict
    inputs: dict


def text_width(text: str, size: float, bold: bool) -> float:
    units = 0.0
    for char in text:
        if unicodedata.combining(char):
            continue
        if char == "\t":
            units += 2.4
        elif unicodedata.east_asian_width(char) in {"F", "W"}:
            units += 1.0
        elif char in "WM@%":
            units += .9
        else:
            units += .58
    return units * size * (1.06 if bold else 1.0)


def check_text(content: str, limits: dict, shape, default_font: float, *, cell: bool = False,
               width_cm: float | None = None, height_cm: float | None = None) -> list:
    parsed = paragraphs(content)
    chars = sum(len(run.text) for p in parsed for run in p.runs)
    max_chars = limits["max_cell_chars" if cell else "max_chars"]
    max_lines = limits["max_cell_lines" if cell else "max_lines"]
    if chars > max_chars:
        raise PptError(f"字符数 {chars} 超过限制 {max_chars}")
    tf = shape.text_frame if not cell else shape
    w_pt = max(1.0, ((width_cm * EMU_CM if width_cm else shape.width) - tf.margin_left - tf.margin_right) / 12700)
    h_pt = max(1.0, ((height_cm * EMU_CM if height_cm else shape.height) - tf.margin_top - tf.margin_bottom) / 12700)
    total_lines = 0
    estimated_height = 0.0
    for paragraph in parsed:
        largest = default_font
        width = 0.0
        for run in paragraph.runs:
            size = run.style.get("font-size", default_font)
            if size > limits["max_font_pt"]:
                raise PptError(f"字号 {size} pt 超过限制 {limits['max_font_pt']} pt")
            largest = max(largest, size)
            width += text_width(run.text, size, run.style.get("font-weight", False))
        if paragraph.marker is None:
            available = w_pt
        else:
            indent = (paragraph.level + 1) * largest
            if paragraph.marker[0] == "num" and paragraph.marker[2] >= 10:
                indent += largest * .6
            available = max(1.0, w_pt - indent)
        line_count = max(1, math.ceil(width / available))
        total_lines += line_count
        estimated_height += line_count * largest * 1.2
    if total_lines > max_lines:
        raise PptError(f"估算行数 {total_lines} 超过限制 {max_lines}，请精简内容或拆页")
    if estimated_height > h_pt + 1:
        raise PptError(f"估算文本高度 {estimated_height / 28.3465:.2f} cm 超过输入框 {h_pt / 28.3465:.2f} cm")
    return parsed


def dimensions(value: dict, limits: dict) -> tuple[float, float]:
    width, height = value.get("width_cm", limits["max_width_cm"]), value.get("height_cm", limits["max_height_cm"])
    if width > limits["max_width_cm"] + .0001 or height > limits["max_height_cm"] + .0001:
        raise PptError(f"宽高 {width} × {height} cm 超过限制 {limits['max_width_cm']} × {limits['max_height_cm']} cm")
    return width, height


def prepare_image(value: dict, base: Path, limits: dict):
    path = local_path(value["src"], base)
    if not path.is_file():
        raise PptError(f"本地图片不存在：{path}")
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            with Image.open(path) as source:
                if source.width * source.height > limits["max_pixels"]:
                    raise PptError(f"图片像素数超过 {limits['max_pixels']}")
                source.load()
                image = ImageOps.exif_transpose(source)
                stream = BytesIO()
                # Normalize orientation and format locally. Keep alpha when present.
                image.convert("RGBA" if image.mode in {"RGBA", "LA", "P"} else "RGB").save(stream, "PNG")
                return stream.getvalue(), image.size
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError, Image.DecompressionBombWarning) as exc:
        raise PptError(f"图片无法读取：{path.name}：{exc}") from exc


def prepare_array(prs, layout: dict, slot: dict, content: str, base: Path) -> list:
    """校验并展开数组条目；返回值按条目次序供 fill_array 使用。"""
    items = array_content(content)
    limits = slot["limits"]
    minimum = limits.get("min_items", 1)
    if not minimum <= len(items) <= limits["max_items"]:
        raise PptError(f"条目数 {len(items)} 超出范围 {minimum}..{limits['max_items']}")
    box, step = slot["box_cm"], slot["step_cm"]
    width, height = prs.slide_width / EMU_CM, prs.slide_height / EMU_CM
    for i in range(len(items)):
        x, y = box["x_cm"] + step["x_cm"] * i, box["y_cm"] + step["y_cm"] * i
        if x < -.001 or y < -.001 or x + box["width_cm"] > width + .001 or y + box["height_cm"] > height + .001:
            raise PptError(f"按样例间距排列第 {i + 1} 个条目将超出页面，请减少条目个数")
    if slot["item_kind"] == "image":
        image_limits = {"max_width_cm": box["width_cm"], "max_height_cm": box["height_cm"],
                        "max_pixels": limits["max_item_pixels"]}
        prepared = []
        for item in items:
            if isinstance(item, list):
                raise PptError("图片数组条目须为路径或图片 JSON 对象")
            value = image_value(item if isinstance(item, dict) else {"src": item})
            value["image_bytes"], value["pixel_size"] = prepare_image(value, base, image_limits)
            prepared.append(value)
        return prepared
    mold = get_shape(prs.slides[layout["template_slide"] - 1], slot["shape_ids"][0])
    if slot["item_kind"] == "group":
        descendants = inner_shapes(mold)
        field_shapes = [descendants[i] for i in slot["field_indexes"]]
    else:
        field_shapes = [mold]
    item_limits = {"max_chars": limits["max_item_chars"], "max_lines": limits["max_item_lines"],
                   "max_font_pt": limits["max_item_font_pt"]}
    fonts = [effective_font(shape) for shape in field_shapes]
    prepared = []
    for item in items:
        fields = [item] if isinstance(item, str) else item
        if not isinstance(fields, list):
            raise PptError("文本/组合数组条目须为字符串或字段列表")
        if len(fields) > len(field_shapes):
            raise PptError(f"条目字段数 {len(fields)} 超过条目的 {len(field_shapes)} 个文本位")
        prepared.append([check_text(field, item_limits, field_shapes[j], fonts[j])
                         for j, field in enumerate(fields)])
    return prepared


def prepare(template: Path, content: Path):
    manifest, data = load_template(template)
    pages = parse_pages(content.read_text(encoding="utf-8-sig"))
    prs = Presentation(BytesIO(data))
    layouts = {layout["id"]: layout for layout in manifest["layouts"]}
    prepared = []
    for page in pages:
        if page.layout not in layouts:
            raise PptError(f"页面 {page.id}：未知布局 {page.layout}；可用：{', '.join(layouts)}")
        layout = layouts[page.layout]
        slots = {slot["id"]: slot for slot in layout["inputs"]}
        extras = set(page.elements) - set(slots)
        if extras:
            raise PptError(f"页面 {page.id}：未知输入元素 {', '.join(sorted(extras))}")
        inputs = {}
        for sid, slot in slots.items():
            if sid not in page.elements:
                # A template input may be omitted from the Markdown.  Missing
                # inputs are rendered as empty content (the output stage
                # removes the corresponding input shape), regardless of the
                # manifest's advisory `required` flag.
                continue
            element = page.elements[sid]
            try:
                if element.kind != slot["type"]:
                    raise PptError(f"输入类型应为 {slot['type']}，实际为 {element.kind}")
                if element.kind == "array":
                    inputs[sid] = prepare_array(prs, layout, slot, element.content, content.parent)
                    continue
                shape = get_shape(prs.slides[layout["template_slide"] - 1], slot["shape_id"])
                if element.kind == "text":
                    inputs[sid] = check_text(element.content, slot["limits"], shape, slot["default_font_pt"],
                                             width_cm=slot["limits"]["max_width_cm"], height_cm=slot["limits"]["max_height_cm"])
                elif element.kind == "image":
                    value = image_content(element.content)
                    value["dimensions"] = dimensions(value, slot["limits"])
                    value["image_bytes"], value["pixel_size"] = prepare_image(value, content.parent, slot["limits"])
                    inputs[sid] = value
                else:
                    value = table_content(element.content)
                    rows, cols = len(value["rows"]), len(value["rows"][0])
                    limits = slot["limits"]
                    if rows > limits["max_rows"] or cols > limits["max_cols"]:
                        raise PptError(f"表格 {rows} 行 × {cols} 列超过 {limits['max_rows']} 行 × {limits['max_cols']} 列（含表头）")
                    width, height = dimensions(value, limits)
                    widths = value.get("column_widths_cm", [width / cols] * cols)
                    if abs(sum(widths) - width) > .001:
                        raise PptError("column_widths_cm 总和必须等于表格 width_cm")
                    value["dimensions"], value["widths"] = (width, height), widths
                    value["parsed_rows"] = [[check_text(cell, limits, shape.table.cell(r, c).text_frame,
                                                        slot["default_font_pt"], cell=True,
                                                        width_cm=widths[c], height_cm=height / rows)
                                             for c, cell in enumerate(row)] for r, row in enumerate(value["rows"])]
                    inputs[sid] = value
            except PptError as exc:
                raise PptError(f"页面 {page.id}/{sid}（MD 第 {element.line} 行）：{exc}") from exc
        prepared.append(Prepared(page.id, layout, inputs))
    return data, prepared


def apply_style(run, style: dict):
    font = run.font
    if "font-weight" in style:
        font.bold = style["font-weight"]
    if "font-style" in style:
        font.italic = style["font-style"]
    if "text-decoration" in style:
        font.underline = style["text-decoration"]
    if "font-size" in style:
        font.size = Pt(style["font-size"])
    if "color" in style:
        font.color.rgb = RGBColor.from_string(style["color"])
    if "font-family" in style:
        font.name = style["font-family"]
        rpr = run._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            old = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag.split(":")[1])
            if old is not None:
                rpr.remove(old)
            child = OxmlElement(tag)
            child.set("typeface", style["font-family"])
            rpr.append(child)


def marker_elements(marker) -> list:
    """Paragraph marker → OOXML bullet elements in a:pPr schema order."""
    if marker is None:
        return [OxmlElement("a:buNone")]
    if marker[0] == "num":
        element = OxmlElement("a:buAutoNum")
        element.set("type", marker[1])
        element.set("startAt", str(marker[2]))
        return [element]
    _, glyph, color, font = marker
    elements = []
    if color:
        bu_clr = OxmlElement("a:buClr")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", color)
        bu_clr.append(srgb)
        elements.append(bu_clr)
    if font:
        bu_font = OxmlElement("a:buFont")
        bu_font.set("typeface", font)
        elements.append(bu_font)
    char = OxmlElement("a:buChar")
    char.set("char", glyph)
    elements.append(char)
    return elements


def fill_text(tf, parsed: list):
    prototype = tf.paragraphs[0]
    ppr = deepcopy(prototype._p.pPr) if prototype._p.pPr is not None else None
    rpr = deepcopy(prototype.runs[0]._r.rPr) if prototype.runs and prototype.runs[0]._r.rPr is not None else None
    tf.clear()
    tf.word_wrap = True
    for index, spec in enumerate(parsed):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        current = paragraph._p.pPr
        if current is not None:
            paragraph._p.remove(current)
        if ppr is not None:
            paragraph._p.insert(0, deepcopy(ppr))
        pp = paragraph._p.get_or_add_pPr()
        # Explicitly choose bullet/no-bullet so master bullets don't leak into plain text.
        for child in list(pp):
            if child.tag.endswith(("}buClr", "}buClrTx", "}buSzTx", "}buSzPct", "}buSzPts", "}buFont",
                                   "}buFontTx", "}buNone", "}buChar", "}buAutoNum", "}buBlip")):
                pp.remove(child)
        if spec.marker is not None:
            paragraph.level = spec.level
            pp.set("marL", str(342900 * (spec.level + 1)))
            pp.set("indent", "-342900")
        insert_at = next((i for i, node in enumerate(pp)
                          if node.tag.endswith(("}tabLst", "}defRPr", "}extLst"))), len(pp))
        for bullet in marker_elements(spec.marker):
            pp.insert(insert_at, bullet)
            insert_at += 1
        for item in spec.runs:
            run = paragraph.add_run()
            if rpr is not None:
                run._r.insert(0, deepcopy(rpr))
            run.text = item.text
            apply_style(run, item.style)


def fill_image(slide, shape, value):
    width, height = (int(Cm(size)) for size in value["dimensions"])
    left, top = shape.left + (shape.width - width) // 2, shape.top + (shape.height - height) // 2
    stream = BytesIO(value["image_bytes"])
    if value.get("fit", "contain") == "contain":
        pw, ph = value["pixel_size"]
        factor = min(width / pw, height / ph)
        iw, ih = round(pw * factor), round(ph * factor)
        new = slide.shapes.add_picture(stream, left + (width - iw) // 2, top + (height - ih) // 2, iw, ih)
    else:
        new = slide.shapes.add_picture(stream, left, top, width, height)
        pw, ph = value["pixel_size"]
        image_ratio, box_ratio = pw / ph, width / height
        if image_ratio > box_ratio:
            new.crop_left = new.crop_right = (1 - box_ratio / image_ratio) / 2
        else:
            new.crop_top = new.crop_bottom = (1 - image_ratio / box_ratio) / 2
    # Transfer image frame effects, rotation and stable ID/name to the replacement.
    old_props = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
    new_props = new._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}spPr")
    if old_props is not None:
        for child in list(old_props):
            if not child.tag.endswith("}xfrm"):
                for existing in list(new_props):
                    if existing.tag == child.tag:
                        new_props.remove(existing)
                new_props.append(deepcopy(child))
    new.rotation = shape.rotation
    c_nv = new._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    old_nv = shape._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
    for key, val in old_nv.attrib.items():
        c_nv.set(key, val)
    if value.get("alt"):
        c_nv.set("descr", value["alt"])
    tree = shape._element.getparent()
    tree.remove(new._element)
    tree.replace(shape._element, new._element)
    # Drop placeholder image relations that are no longer used.
    used = {val for node in slide._element.iter() for attr, val in node.attrib.items()
            if attr.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}")}
    for rel in list(slide.part.rels.values()):
        if rel.reltype.endswith("/image") and rel.rId not in used:
            slide.part.drop_rel(rel.rId)


def fill_table(shape, value):
    table = shape.table
    rows, cols = len(value["parsed_rows"]), len(value["parsed_rows"][0])
    # Remove unused cells, grid columns and rows, preserving native table styles.
    for row in list(table._tbl.tr_lst)[rows:]:
        table._tbl.remove(row)
    for row in table._tbl.tr_lst:
        for cell in list(row.tc_lst)[cols:]:
            row.remove(cell)
    for col in list(table._tbl.tblGrid.gridCol_lst)[cols:]:
        table._tbl.tblGrid.remove(col)
    width, height = value["dimensions"]
    shape.width, shape.height = Cm(width), Cm(height)
    for c, w in enumerate(value["widths"]):
        table.columns[c].width = Cm(w)
    for r in range(rows):
        table.rows[r].height = Cm(height / rows)
        for c in range(cols):
            fill_text(table.cell(r, c).text_frame, value["parsed_rows"][r][c])


def renumber_shape(slide, element) -> int:
    """Give a cloned shape tree fresh cNvPr ids; return the new root id."""
    used = {int(node.get("id")) for node in slide._element.iter()
            if node.tag.endswith("}cNvPr") and node.get("id", "").isdigit()}
    next_id = max(used or {0}) + 1
    root = None
    for node in element.iter():
        if node.tag.endswith("}cNvPr"):
            if root is None:
                root = next_id
            node.set("id", str(next_id))
            next_id += 1
    return root


def fill_array_image(slide, shape, value):
    """Swap a mold/clone picture's blip and adapt it to its own frame."""
    _, rid = slide.part.get_or_add_image_part(BytesIO(value["image_bytes"]))
    blip = shape._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    blip.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed", rid)
    shape.crop_left = shape.crop_right = shape.crop_top = shape.crop_bottom = 0
    width, height = shape.width, shape.height
    pw, ph = value["pixel_size"]
    if value.get("fit", "contain") == "contain":
        factor = min(width / pw, height / ph)
        iw, ih = round(pw * factor), round(ph * factor)
        shape.left += (width - iw) // 2
        shape.top += (height - ih) // 2
        shape.width, shape.height = iw, ih
    else:
        image_ratio, box_ratio = pw / ph, width / height
        if image_ratio > box_ratio:
            shape.crop_left = shape.crop_right = (1 - box_ratio / image_ratio) / 2
        else:
            shape.crop_top = shape.crop_bottom = (1 - image_ratio / box_ratio) / 2


def fill_array(slide, slot: dict, items):
    """数组输入：沿两个样例的间距向量填充/复制条目。"""
    molds = [get_shape(slide, i) for i in slot["shape_ids"]]
    tree = molds[0]._element.getparent()
    if items is None:
        for mold in molds:
            tree.remove(mold._element)
        return
    box, step = slot["box_cm"], slot["step_cm"]
    for i, item in enumerate(items):
        if i < 2:
            target = molds[i]
        else:
            element = deepcopy(molds[1]._element)
            root_id = renumber_shape(slide, element)
            tree.append(element)
            target = get_shape(slide, root_id)
            target.left = int(Cm(box["x_cm"] + step["x_cm"] * i))
            target.top = int(Cm(box["y_cm"] + step["y_cm"] * i))
        if slot["item_kind"] == "image":
            fill_array_image(slide, target, item)
        elif slot["item_kind"] == "group":
            descendants = inner_shapes(target)
            for j, parsed in enumerate(item):
                fill_text(descendants[slot["field_indexes"][j]].text_frame, parsed)
        else:
            fill_text(target.text_frame, item[0])
    if len(items) == 1:
        tree.remove(molds[1]._element)


def drop_unused_images(slide):
    used = {val for node in slide._element.iter() for attr, val in node.attrib.items()
            if attr.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}")}
    for rel in list(slide.part.rels.values()):
        if rel.reltype.endswith("/image") and rel.rId not in used:
            slide.part.drop_rel(rel.rId)


def generate(template: Path, content: Path, output: Path, *, overwrite=False) -> dict:
    # Validate the entire deck before any output mutation.
    data, prepared = prepare(template, content)
    if output == content or output == template / "template.pptx":
        raise PptError("输出不能覆盖模板或填充内容")
    if output.suffix.lower() != ".pptx":
        raise PptError("输出文件扩展名必须为 .pptx")
    prs = Presentation(BytesIO(clone_slides(data, [p.layout["template_slide"] for p in prepared])))
    for slide, page in zip(prs.slides, prepared):
        for slot in page.layout["inputs"]:
            if slot["type"] == "array":
                fill_array(slide, slot, page.inputs.get(slot["id"]))
                continue
            shape = get_shape(slide, slot["shape_id"])
            if slot["id"] not in page.inputs:
                shape._element.getparent().remove(shape._element)
                continue
            value = page.inputs[slot["id"]]
            if slot["type"] == "text":
                shape.width = Cm(slot["limits"]["max_width_cm"])
                shape.height = Cm(slot["limits"]["max_height_cm"])
                fill_text(shape.text_frame, value)
            elif slot["type"] == "image":
                fill_image(slide, shape, value)
            else:
                fill_table(shape, value)
        drop_unused_images(slide)
    buffer = BytesIO()
    prs.save(buffer)
    result = buffer.getvalue()
    inspect_package(result)
    # Verify exported slide count and loadability before writing a final file.
    if len(Presentation(BytesIO(result)).slides) != len(prepared):
        raise PptError("输出页面数量校验失败")
    atomic_write(output, result, overwrite=overwrite)
    return {"output": str(output), "slides": len(prepared), "validation": "passed"}
